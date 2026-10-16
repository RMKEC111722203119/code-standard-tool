import json,math
from collections import Counter
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill,Alignment
from .models import ReportPart
from .utils import debug
COLS=["run_id","application_repository_url","application_branch","application_commit","checks_branch","checks_commit","test_id","test_name","test_version","severity","confidence","category","language","file_path","line_number","column_number","end_line_number","end_column_number","cell_number","symbol","symbol_type","evidence","reason","expected","suggestion","documentation_reference","blocking"]
def findings(path):
    with path.open() as f:
        for line in f:
            if line.strip(): yield json.loads(line)
def enrich(x,run,r,ac,cc):
    z={"run_id":run,"application_repository_url":r.application_repository_url,"application_branch":r.application_branch,"application_commit":ac,"checks_branch":r.checks_branch,"checks_commit":cc,**x}; z["severity"]=str(z.get("severity","Medium")).title(); z["blocking"]=z["severity"] in r.blocking_severities; return z
def summarize(path,r):
    debug(f"finding summary started path={path}")
    c=Counter(); total=block=0
    for x in findings(path):
        s=str(x.get("severity","Medium")).title(); c[s]+=1; total+=1; block+=s in r.blocking_severities
    for s in ["Critical","High","Medium","Low"]: c.setdefault(s,0)
    debug(f"finding summary completed total={total} blocking={block}")
    return total,dict(c),block
def header(ws,cols):
    row=[]
    for v in cols:
        c=WriteOnlyCell(ws,value=v); c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78"); row.append(c)
    ws.append(row)
def excel_parts(dir,path,total,run,r,ac,cc,execs,summary,rows_per_sheet,per_file,errors,warnings):
    debug(f"excel generation started run_id={run} findings={total}")
    dir.mkdir(); out=[]; source=findings(path); idx=0; count=max(1,math.ceil(total/max(1,per_file))); rows_per_sheet=min(rows_per_sheet,1048575)
    for part in range(1,count+1):
        wb=Workbook(write_only=True)
        if part==1:
            ws=wb.create_sheet("Run Summary"); header(ws,["Field","Value"])
            for k,v in summary.items(): ws.append([k,json.dumps(v) if isinstance(v,(dict,list)) else v])
            ws=wb.create_sheet("Test Execution"); fields=list(execs[0].model_dump()) if execs else ["test_id","status"]; header(ws,fields)
            for e in execs: ws.append([json.dumps(v) if isinstance(v,(dict,list)) else v for v in e.model_dump().values()])
        target=min(per_file,total-idx); written=0; sheet=0; ws=None
        while written<target:
            if written%rows_per_sheet==0: sheet+=1; ws=wb.create_sheet(f"Findings_{sheet:03d}"); header(ws,COLS)
            x=enrich(next(source),run,r,ac,cc); ws.append([json.dumps(x.get(c)) if isinstance(x.get(c),(dict,list)) else x.get(c) for c in COLS]); written+=1; idx+=1
        if total==0: ws=wb.create_sheet("Findings_001"); header(ws,COLS)
        name=f"code_standards_{run}_part_{part:03d}.xlsx" if count>1 else f"code_standards_{run}.xlsx"; wb.save(dir/name); out.append(ReportPart(part=part,kind="excel",blob_name="",filename=name,first_finding=(idx-written+1 if total else None),last_finding=(idx if total else None)))
    debug(f"excel generation completed run_id={run} parts={len(out)}")
    return out
def json_parts(dir,path,total,run,r,ac,cc,per_file):
    debug(f"json generation started run_id={run} findings={total}")
    out=[]; source=findings(path); idx=0; count=max(1,math.ceil(total/max(1,per_file)))
    for part in range(1,count+1):
        target=min(per_file,total-idx); name=f"results_{part:03d}.json"; first=idx+1 if total else None
        with (dir/name).open("w") as f:
            f.write("[\n")
            for i in range(target):
                if i:f.write(",\n")
                json.dump(enrich(next(source),run,r,ac,cc),f); idx+=1
            f.write("\n]")
        out.append(ReportPart(part=part,kind="json",blob_name="",filename=name,first_finding=first,last_finding=idx if total else None))
    debug(f"json generation completed run_id={run} parts={len(out)}")
    return out
def write_index(dir,run,repo,branch,summary,excel,jsons):
    debug(f"index write started run_id={run}")
    (dir/"report_index.json").write_text(json.dumps({"run_id":run,"repository_name":repo,"branch_path":branch,"summary":summary,"excel_reports":[x.model_dump() for x in excel],"json_reports":[x.model_dump() for x in jsons]},indent=2))
    debug(f"index write completed run_id={run}")
